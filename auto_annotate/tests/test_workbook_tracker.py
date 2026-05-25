from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from auto_annotate.cli import main
from auto_annotate.workbook_tracker import (
    SHEET_HEADERS,
    check_3dgait_outputs,
    check_visfrailty_outputs,
    init_workbook,
    scan_workbook,
    workbook_status,
)


class WorkbookTrackerTests(unittest.TestCase):
    def test_init_workbook_creates_all_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "frailscreen_flywheel.xlsx"
            init_workbook(path)
            wb = load_workbook(path)

        for sheet_name, headers in SHEET_HEADERS.items():
            self.assertIn(sheet_name, wb.sheetnames)
            self.assertEqual([cell.value for cell in wb[sheet_name][1]][: len(headers)], headers)

    def test_scan_check_and_status_track_current_lifecycle(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_root = root / "input"
            test_folder = input_root / "patient01 TUG"
            test_folder.mkdir(parents=True)
            (test_folder / "rgb_video_sample.mp4").write_bytes(b"video")
            (test_folder / "biomarker.json").write_text(
                json.dumps({"tug": {"start_ms": 1000, "end_ms": 9000, "duration_ms": 8000}}),
                encoding="utf-8",
            )
            (test_folder / "left_boundaries.json").write_text(
                json.dumps({"events": [{"time_ms": 2000, "event_type": "stance_start"}]}),
                encoding="utf-8",
            )
            (test_folder / "right_boundaries.json").write_text(
                json.dumps({"events": [{"time_ms": 2400, "event_type": "stance_start"}]}),
                encoding="utf-8",
            )
            output = root / "output"
            output.mkdir()
            (output / "patient01_TUG_auto.eaf").write_text("<ANNOTATION_DOCUMENT />", encoding="utf-8")
            (output / "patient01_TUG_auto.json").write_text(
                json.dumps({"test_type": "tug", "occlusion_count": 0}),
                encoding="utf-8",
            )
            gt = root / "ground_truth"
            gt.mkdir()
            (gt / "patient01_TUG_ground_truth.eaf").write_text("<ANNOTATION_DOCUMENT />", encoding="utf-8")
            excel = root / "frailscreen_flywheel.xlsx"

            scan = scan_workbook(excel, input_root, output, root / "batch_annotations", gt)
            check = check_3dgait_outputs(excel)
            summary = workbook_status(excel)

        self.assertEqual(scan["videos_found"], 1)
        self.assertEqual(scan["auto_annotations_found"], 1)
        self.assertEqual(scan["human_corrections_found"], 1)
        self.assertEqual(check["videos_checked"], 1)
        self.assertEqual(summary.videos, 1)
        self.assertEqual(summary.pipeline_complete, 1)
        self.assertEqual(summary.auto_annotated, 1)
        self.assertEqual(summary.human_corrected, 1)

    def test_workbook_cli_commands(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_root = root / "input" / "sample FT"
            input_root.mkdir(parents=True)
            (input_root / "rgb_video_sample.mp4").write_bytes(b"video")
            excel = root / "flywheel.xlsx"

            self.assertEqual(main(["workbook-init", "--excel", str(excel)]), 0)
            self.assertEqual(
                main([
                    "workbook-scan",
                    "--excel",
                    str(excel),
                    "--input",
                    str(root / "input"),
                    "--output",
                    str(root / "output"),
                    "--batch-output",
                    str(root / "batch_annotations"),
                    "--ground-truth",
                    str(root / "ground_truth"),
                ]),
                0,
            )
            self.assertEqual(main(["check-3dgait", "--excel", str(excel)]), 0)
            self.assertEqual(main(["check-visfrailty", "--excel", str(excel)]), 0)
            self.assertEqual(main(["workbook-status", "--excel", str(excel), "--json"]), 0)

    def test_visfrailty_v7_workbook_syncs_annotation_flywheel(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            excel = root / "visfrailty_screening_v7.xlsx"
            output = root / "output"
            output.mkdir()
            gt = root / "ground_truth"
            gt.mkdir()
            pipeline_run = root / "visfrailty_outputs" / "20260501_MC0001"
            test_out = pipeline_run / "MC0001 TUG1"
            test_out.mkdir(parents=True)
            (test_out / "sppb_results.json").write_text("{}", encoding="utf-8")
            (output / "MC0001_TUG1_auto.eaf").write_text("<ANNOTATION_DOCUMENT />", encoding="utf-8")
            (output / "MC0001_TUG1_auto.json").write_text(
                json.dumps({"test_type": "tug", "phase_unknown_count": 1, "foot_unknown_count": 2}),
                encoding="utf-8",
            )
            (gt / "MC0001_TUG1_ground_truth.eaf").write_text("<ANNOTATION_DOCUMENT />", encoding="utf-8")
            _write_minimal_v7_workbook(excel, pipeline_run)

            check = check_visfrailty_outputs(excel)
            scan = scan_workbook(excel, root / "input", output, root / "batch_annotations", gt)
            summary = workbook_status(excel)
            wb = load_workbook(excel, data_only=True)
            ws = wb["annotation_flywheel"]
            headers = {cell.value: idx for idx, cell in enumerate(ws[3], start=1)}

        self.assertEqual(check["videos_checked"], 1)
        self.assertEqual(scan["videos_found"], 1)
        self.assertEqual(scan["auto_annotations_found"], 1)
        self.assertEqual(scan["human_corrections_found"], 1)
        self.assertEqual(ws.cell(row=4, column=headers["pipeline_state"]).value, "processed")
        self.assertEqual(ws.cell(row=4, column=headers["auto_eaf_generated"]).value, "yes")
        self.assertEqual(ws.cell(row=4, column=headers["auto_num_unknowns"]).value, 3)
        self.assertEqual(ws.cell(row=4, column=headers["human_review_status"]).value, "corrected")
        self.assertEqual(ws.cell(row=4, column=headers["flywheel_stage"]).value, "ground_truth_confirmed")
        self.assertEqual(summary.videos, 1)
        self.assertEqual(summary.pipeline_complete, 1)
        self.assertEqual(summary.auto_annotated, 1)
        self.assertEqual(summary.human_corrected, 1)


def _write_minimal_v7_workbook(path: Path, pipeline_run: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "annotation_flywheel"
    ws.append(["FRAILScreen - Annotation Flywheel"])
    ws.append(["Identity", None, None, None, "Pipeline", None, None, "Annotation"])
    ws.append(
        [
            "subject_id",
            "site",
            "video_slot",
            "slot_type",
            "pipeline_state",
            "biomarker_json",
            "pipeline_error",
            "flywheel_stage",
            "auto_eaf_generated",
            "auto_eaf_path",
            "auto_num_unknowns",
            "auto_correction_min",
            "human_review_status",
            "human_reviewer",
            "human_review_date",
            "human_eaf_path",
            "ground_truth_confirmed",
            "ground_truth_date",
            "usable_for_training",
            "exclusion_reason",
        ]
    )
    ws.append(["MC0001", "MC", "TUG1", "tug", "missing", None, None, "raw", None, None, None, None, "not_started"])
    dashboard = wb.create_sheet("flywheel_dashboard")
    for _ in range(10):
        dashboard.append([None, None, None])
    for row, stage in enumerate(["raw", "pipeline_processed", "auto_annotated", "human_reviewed", "ground_truth_confirmed"], start=5):
        dashboard.cell(row=row, column=1).value = stage
    subjects = wb.create_sheet("subjects")
    subjects.append(["subject_id", "output_run_path", "processing_state"])
    subjects.append(["MC0001", str(pipeline_run), "raw"])
    wb.save(path)


if __name__ == "__main__":
    unittest.main()
