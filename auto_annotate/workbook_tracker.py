"""Excel workbook tracker for the FrailScreen annotation flywheel."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
import json
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .pipeline_adapter import detect_test_type

VISFRAILTY_V7_ANNOTATION_SHEET = "annotation_flywheel"
VISFRAILTY_V7_DASHBOARD_SHEET = "flywheel_dashboard"
VISFRAILTY_V7_SUBJECTS_SHEET = "subjects"
VISFRAILTY_V7_HEADER_ROW = 3
VISFRAILTY_V7_STAGES = (
    "raw",
    "pipeline_processed",
    "auto_annotated",
    "human_reviewed",
    "ground_truth_confirmed",
)
_VISFRAILTY_PATIENT_OUTPUT_CACHE: dict[tuple[str, str], list[Path]] = {}

SHEET_HEADERS: dict[str, list[str]] = {
    "Videos": [
        "video_id",
        "patient_id",
        "test_type",
        "video_path",
        "input_folder",
        "recording_date",
        "duration_ms",
        "split",
        "status",
        "notes",
    ],
    "PipelineOutputs": [
        "video_id",
        "pipeline_out_folder",
        "biomarker_json",
        "sppb_results_json",
        "gait_boundaries_left",
        "gait_boundaries_right",
        "tug_timing_found",
        "gait_events_found",
        "pipeline_status",
        "last_checked_at",
        "error_message",
    ],
    "AutoAnnotations": [
        "auto_run_id",
        "video_id",
        "test_type",
        "config_id",
        "code_version",
        "auto_eaf_path",
        "debug_json_path",
        "run_timestamp",
        "phase_unknown_count",
        "foot_unknown_count",
        "occlusion_count",
        "validation_status",
        "eaf_written",
        "exit_code",
        "error_message",
    ],
    "HumanCorrections": [
        "correction_id",
        "video_id",
        "auto_run_id",
        "reviewer",
        "auto_eaf_path",
        "ground_truth_eaf_path",
        "corrected_at",
        "review_time_minutes",
        "correction_status",
        "notes",
    ],
    "GroundTruthValidation": [
        "video_id",
        "ground_truth_eaf_path",
        "validated_at",
        "valid",
        "test_tier_ok",
        "phase_tier_ok",
        "foot_tiers_ok",
        "occlusion_tier_ok",
        "no_out_of_bounds_annotations",
        "error_message",
    ],
    "CalibrationRuns": [
        "calibration_id",
        "started_at",
        "finished_at",
        "config_in",
        "config_out",
        "train_count",
        "validation_count",
        "holdout_count",
        "overall_score",
        "phase_score",
        "foot_score",
        "test_score",
        "occlusion_score",
        "accepted",
        "notes",
    ],
    "CalibrationSamples": [
        "calibration_id",
        "video_id",
        "test_type",
        "role",
        "included",
        "reason_excluded",
    ],
    "MetricsHistory": [
        "timestamp",
        "config_id",
        "test_type",
        "split",
        "sample_count",
        "test_interval_score",
        "phase_overlap_score",
        "phase_boundary_mae_ms",
        "left_foot_score",
        "right_foot_score",
        "occlusion_score",
        "unknown_rate",
    ],
    "ReviewQueue": [
        "priority",
        "video_id",
        "test_type",
        "auto_eaf_path",
        "reason",
        "unknown_count",
        "validation_status",
        "suggested_action",
    ],
    "Configs": [
        "config_id",
        "config_path",
        "created_at",
        "created_by",
        "source_calibration_id",
        "description",
        "active",
    ],
    "EdgeCases": [
        "video_id",
        "test_type",
        "edge_case_type",
        "severity",
        "notes",
    ],
}

TRACKED_SHEETS = tuple(SHEET_HEADERS)
LIST_SHEET = "_Lists"
YES_NO = ("TRUE", "FALSE")
TEST_TYPES = ("TUG", "GS1", "GS2", "CRT", "SBS", "ST", "FT", "unknown")
SPLITS = ("unassigned", "train", "validation", "holdout")
VIDEO_STATUSES = (
    "registered",
    "visfrailty_missing",
    "visfrailty_partial",
    "visfrailty_complete",
    "auto_annotated",
    "human_corrected",
    "validated",
    "used_for_calibration",
)
PIPELINE_STATUSES = ("missing", "partial", "complete", "failed")
VALIDATION_STATUSES = ("unknown", "valid", "invalid")
CORRECTION_STATUSES = ("pending", "corrected", "rejected", "needs_second_review")
EDGE_CASE_TYPES = (
    "pivot_turner",
    "bystander",
    "occluded_feet",
    "partial_test",
    "slow_sit_to_stand",
    "poor_gait_boundaries",
    "patient_out_of_frame",
)


@dataclass(frozen=True)
class WorkbookSummary:
    videos: int
    pipeline_complete: int
    pipeline_partial: int
    pipeline_missing: int
    auto_annotated: int
    human_corrected: int
    ground_truth_validated: int
    calibration_sample_videos: int
    ready_for_auto_annotation: int
    ready_for_human_review: int
    ready_for_calibration: int
    review_queue: int

    def to_dict(self) -> dict[str, int]:
        return self.__dict__.copy()


def init_workbook(excel_path: str | Path) -> Path:
    """Create or upgrade the Excel flywheel workbook."""

    path = Path(excel_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and _is_visfrailty_v7_workbook(path):
        return path
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
    _ensure_workbook_shape(wb)
    wb.save(path)
    return path


def scan_workbook(
    excel_path: str | Path,
    input_root: str | Path,
    output_root: str | Path = "output",
    batch_output_root: str | Path = "batch_annotations",
    ground_truth_root: str | Path = "ground_truth",
) -> dict[str, int]:
    """Scan local files and update Videos/AutoAnnotations/HumanCorrections."""

    path = Path(excel_path)
    if path.exists() and _is_visfrailty_v7_workbook(path):
        return sync_visfrailty_v7_workbook(
            path,
            output_root=output_root,
            batch_output_root=batch_output_root,
            ground_truth_root=ground_truth_root,
        )

    path = init_workbook(path)
    wb = load_workbook(path)
    input_path = Path(input_root)
    output_path = Path(output_root)
    batch_output_path = Path(batch_output_root)
    ground_truth_path = Path(ground_truth_root)

    videos = _scan_video_folders(input_path)
    existing_video_rows = _sheet_rows_by_key(wb["Videos"], "video_id")
    for video in videos:
        row = existing_video_rows.get(video["video_id"], {})
        row.update({key: value for key, value in video.items() if value is not None})
        row.setdefault("split", "unassigned")
        row.setdefault("status", "registered")
        existing_video_rows[video["video_id"]] = row
    _write_keyed_rows(wb["Videos"], "video_id", existing_video_rows)

    auto_rows = _auto_annotation_rows(output_path, batch_output_path)
    _write_keyed_rows(wb["AutoAnnotations"], "auto_run_id", auto_rows)

    correction_rows = _human_correction_rows(ground_truth_path, existing_video_rows)
    _write_keyed_rows(wb["HumanCorrections"], "correction_id", correction_rows)

    _update_video_statuses(wb)
    _refresh_review_queue(wb)
    _ensure_workbook_shape(wb)
    wb.save(path)
    return {
        "videos_found": len(videos),
        "auto_annotations_found": len(auto_rows),
        "human_corrections_found": len(correction_rows),
    }


def check_visfrailty_outputs(
    excel_path: str | Path,
    visfrailty_root: str | Path | None = None,
) -> dict[str, int]:
    """Update pipeline status from VisFrailTy/visfrailty_screening_v7 outputs."""

    path = Path(excel_path)
    if path.exists() and _is_visfrailty_v7_workbook(path):
        return _check_visfrailty_v7_outputs(path, visfrailty_root)

    path = init_workbook(path)
    wb = load_workbook(path)
    video_rows = list(_iter_rows_as_dicts(wb["Videos"]))
    pipeline_rows = _sheet_rows_by_key(wb["PipelineOutputs"], "video_id")
    checked = 0
    for video in video_rows:
        video_id = str(video.get("video_id") or "")
        if not video_id:
            continue
        input_folder = Path(str(video.get("input_folder") or ""))
        existing = pipeline_rows.get(video_id, {})
        pipeline_folder = _resolve_visfrailty_output_folder(
            input_folder=input_folder,
            video_id=video_id,
            test_type=str(video.get("test_type") or ""),
            existing_folder=Path(str(existing.get("pipeline_out_folder") or "")),
            visfrailty_root=Path(visfrailty_root) if visfrailty_root else None,
        )
        pipeline_rows[video_id] = _pipeline_status_row(video_id, pipeline_folder, str(video.get("test_type") or ""))
        checked += 1

    _write_keyed_rows(wb["PipelineOutputs"], "video_id", pipeline_rows)
    _update_video_statuses(wb)
    _refresh_review_queue(wb)
    _ensure_workbook_shape(wb)
    wb.save(path)
    return {"videos_checked": checked}


def check_3dgait_outputs(
    excel_path: str | Path,
    visfrailty_root: str | Path | None = None,
) -> dict[str, int]:
    """Compatibility alias for older CLI/tests; VisFrailTy is now the pipeline source."""

    return check_visfrailty_outputs(excel_path, visfrailty_root)


def sync_visfrailty_v7_workbook(
    excel_path: str | Path,
    output_root: str | Path = "output",
    batch_output_root: str | Path = "batch_annotations",
    ground_truth_root: str | Path = "ground_truth",
) -> dict[str, int]:
    """Sync auto-annotation/review progress into visfrailty_screening_v7.xlsx."""

    path = Path(excel_path)
    wb = load_workbook(path)
    ws = wb[VISFRAILTY_V7_ANNOTATION_SHEET]
    headers = _v7_header_map(ws, VISFRAILTY_V7_HEADER_ROW)
    required = {"subject_id", "video_slot", "flywheel_stage"}
    if not required.issubset(headers):
        missing = ", ".join(sorted(required - set(headers)))
        raise ValueError(f"{VISFRAILTY_V7_ANNOTATION_SHEET} missing columns: {missing}")

    autos = _v7_auto_annotation_map(Path(output_root), Path(batch_output_root))
    corrections = _v7_ground_truth_map(Path(ground_truth_root))
    updated = 0
    for row_idx in range(VISFRAILTY_V7_HEADER_ROW + 1, ws.max_row + 1):
        subject_id = _cell_value(ws, row_idx, headers, "subject_id")
        video_slot = _cell_value(ws, row_idx, headers, "video_slot")
        if not subject_id or not video_slot:
            continue
        key = (str(subject_id).upper(), str(video_slot).upper())
        auto = autos.get(key)
        correction = corrections.get(key)
        if auto:
            _set_cell(ws, row_idx, headers, "auto_eaf_generated", "yes")
            _set_cell(ws, row_idx, headers, "auto_eaf_path", auto.get("auto_eaf_path", ""))
            unknowns = _v7_unknown_count(auto)
            if unknowns != "":
                _set_cell(ws, row_idx, headers, "auto_num_unknowns", unknowns)
            updated += 1
        if correction:
            _set_cell(ws, row_idx, headers, "human_review_status", "corrected")
            _set_cell(ws, row_idx, headers, "human_eaf_path", correction)
            _set_cell(ws, row_idx, headers, "ground_truth_confirmed", "yes")
            _set_cell(ws, row_idx, headers, "ground_truth_date", _file_mtime_iso(Path(correction)))
            _set_cell(ws, row_idx, headers, "usable_for_training", "yes")
            updated += 1
        _set_cell(ws, row_idx, headers, "flywheel_stage", _v7_stage_for_row(ws, row_idx, headers))

    _refresh_visfrailty_v7_dashboard(wb)
    wb.save(path)
    return {
        "videos_found": _count_v7_annotation_rows(ws),
        "auto_annotations_found": len(autos),
        "human_corrections_found": len(corrections),
        "rows_updated": updated,
    }


def workbook_status(excel_path: str | Path) -> WorkbookSummary:
    """Return current flywheel counts from the Excel workbook."""

    path = Path(excel_path)
    if path.exists() and _is_visfrailty_v7_workbook(path):
        return _visfrailty_v7_status(path)

    path = init_workbook(path)
    wb = load_workbook(path, data_only=True)
    videos = list(_iter_rows_as_dicts(wb["Videos"]))
    pipeline = list(_iter_rows_as_dicts(wb["PipelineOutputs"]))
    autos = list(_iter_rows_as_dicts(wb["AutoAnnotations"]))
    corrections = list(_iter_rows_as_dicts(wb["HumanCorrections"]))
    validations = list(_iter_rows_as_dicts(wb["GroundTruthValidation"]))
    calibration_samples = list(_iter_rows_as_dicts(wb["CalibrationSamples"]))
    review_queue = list(_iter_rows_as_dicts(wb["ReviewQueue"]))

    pipeline_by_video = {row.get("video_id"): row for row in pipeline}
    auto_videos = {row.get("video_id") for row in autos if _truthy(row.get("eaf_written"))}
    corrected_videos = {
        row.get("video_id")
        for row in corrections
        if str(row.get("correction_status") or "").lower() == "corrected"
        and row.get("ground_truth_eaf_path")
    }
    validated_videos = {row.get("video_id") for row in validations if _truthy(row.get("valid"))}
    calibration_videos = {
        row.get("video_id")
        for row in calibration_samples
        if _truthy(row.get("included"))
    }

    pipeline_complete = sum(1 for row in pipeline if row.get("pipeline_status") == "complete")
    pipeline_partial = sum(1 for row in pipeline if row.get("pipeline_status") == "partial")
    pipeline_missing = sum(1 for row in pipeline if row.get("pipeline_status") == "missing")
    ready_for_auto = sum(
        1
        for row in videos
        if (pipeline_by_video.get(row.get("video_id")) or {}).get("pipeline_status") == "complete"
        and row.get("video_id") not in auto_videos
    )
    ready_for_review = sum(1 for video_id in auto_videos if video_id not in corrected_videos)
    valid_uncalibrated = validated_videos - calibration_videos
    corrected_uncalibrated = corrected_videos - calibration_videos
    return WorkbookSummary(
        videos=len(videos),
        pipeline_complete=pipeline_complete,
        pipeline_partial=pipeline_partial,
        pipeline_missing=pipeline_missing,
        auto_annotated=len(auto_videos),
        human_corrected=len(corrected_videos),
        ground_truth_validated=len(validated_videos),
        calibration_sample_videos=len(calibration_videos),
        ready_for_auto_annotation=ready_for_auto,
        ready_for_human_review=ready_for_review,
        ready_for_calibration=len(valid_uncalibrated or corrected_uncalibrated),
        review_queue=len(review_queue),
    )


def _ensure_workbook_shape(wb: Any) -> None:
    for sheet_name, headers in SHEET_HEADERS.items():
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
        _ensure_headers(ws, headers)
        _style_sheet(ws)
    _ensure_lists_sheet(wb)
    _apply_data_validations(wb)


def _ensure_headers(ws: Any, headers: list[str]) -> None:
    existing = [ws.cell(row=1, column=idx + 1).value for idx in range(len(headers))]
    if existing != headers:
        for idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=idx).value = header


def _style_sheet(ws: Any) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for column_cells in ws.columns:
        header = str(column_cells[0].value or "")
        width = max(12, min(60, len(header) + 4))
        ws.column_dimensions[column_cells[0].column_letter].width = width


def _ensure_lists_sheet(wb: Any) -> None:
    ws = wb[LIST_SHEET] if LIST_SHEET in wb.sheetnames else wb.create_sheet(LIST_SHEET)
    lists = {
        "A": TEST_TYPES,
        "B": SPLITS,
        "C": VIDEO_STATUSES,
        "D": PIPELINE_STATUSES,
        "E": VALIDATION_STATUSES,
        "F": CORRECTION_STATUSES,
        "G": YES_NO,
        "H": EDGE_CASE_TYPES,
    }
    headers = {
        "A": "test_types",
        "B": "splits",
        "C": "video_statuses",
        "D": "pipeline_statuses",
        "E": "validation_statuses",
        "F": "correction_statuses",
        "G": "yes_no",
        "H": "edge_case_types",
    }
    for column, values in lists.items():
        ws[f"{column}1"] = headers[column]
        for row_idx, value in enumerate(values, start=2):
            ws[f"{column}{row_idx}"] = value
    ws.sheet_state = "hidden"


def _apply_data_validations(wb: Any) -> None:
    validations = [
        ("Videos", "test_type", f"'{LIST_SHEET}'!$A$2:$A${len(TEST_TYPES) + 1}"),
        ("Videos", "split", f"'{LIST_SHEET}'!$B$2:$B${len(SPLITS) + 1}"),
        ("Videos", "status", f"'{LIST_SHEET}'!$C$2:$C${len(VIDEO_STATUSES) + 1}"),
        ("PipelineOutputs", "pipeline_status", f"'{LIST_SHEET}'!$D$2:$D${len(PIPELINE_STATUSES) + 1}"),
        ("AutoAnnotations", "validation_status", f"'{LIST_SHEET}'!$E$2:$E${len(VALIDATION_STATUSES) + 1}"),
        ("AutoAnnotations", "eaf_written", f"'{LIST_SHEET}'!$G$2:$G${len(YES_NO) + 1}"),
        ("HumanCorrections", "correction_status", f"'{LIST_SHEET}'!$F$2:$F${len(CORRECTION_STATUSES) + 1}"),
        ("GroundTruthValidation", "valid", f"'{LIST_SHEET}'!$G$2:$G${len(YES_NO) + 1}"),
        ("CalibrationSamples", "role", f"'{LIST_SHEET}'!$B$3:$B${len(SPLITS) + 1}"),
        ("CalibrationSamples", "included", f"'{LIST_SHEET}'!$G$2:$G${len(YES_NO) + 1}"),
        ("Configs", "active", f"'{LIST_SHEET}'!$G$2:$G${len(YES_NO) + 1}"),
        ("EdgeCases", "edge_case_type", f"'{LIST_SHEET}'!$H$2:$H${len(EDGE_CASE_TYPES) + 1}"),
    ]
    for sheet_name, header, formula in validations:
        ws = wb[sheet_name]
        col = _column_for_header(ws, header)
        if col is None:
            continue
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(validation)
        validation.add(f"{col}2:{col}5000")


def _scan_video_folders(input_root: Path) -> list[dict[str, Any]]:
    if not input_root.exists():
        return []
    folders = sorted(
        {path.parent for path in input_root.rglob("rgb_video*.mp4")},
        key=lambda path: str(path).lower(),
    )
    used_ids: set[str] = set()
    rows: list[dict[str, Any]] = []
    for folder in folders:
        videos = sorted(folder.glob("rgb_video*.mp4"))
        if not videos:
            continue
        video_path = videos[0]
        relative = _safe_relative(folder, input_root)
        base_id = _safe_stem(str(relative)) if str(relative) != "." else _safe_stem(folder.name)
        video_id = _unique_id(base_id, used_ids)
        rows.append(
            {
                "video_id": video_id,
                "patient_id": _infer_patient_id(folder),
                "test_type": _excel_test_type(detect_test_type(folder)),
                "video_path": str(video_path),
                "input_folder": str(folder),
                "recording_date": "",
                "duration_ms": "",
                "split": "unassigned",
                "status": "registered",
                "notes": "",
            }
        )
    return rows


def _auto_annotation_rows(output_root: Path, batch_output_root: Path) -> dict[str, dict[str, Any]]:
    rows: dict[str, dict[str, Any]] = {}
    for root in (output_root, batch_output_root):
        if not root.exists():
            continue
        for eaf in sorted(root.rglob("*_auto.eaf")):
            meta = _read_json(eaf.with_suffix(".json")) or {}
            auto_run_id = _safe_stem(str(eaf.relative_to(root)))
            video_id = _video_id_from_auto_eaf(eaf)
            rows[auto_run_id] = {
                "auto_run_id": auto_run_id,
                "video_id": video_id,
                "test_type": _excel_test_type(meta.get("test_type")),
                "config_id": "",
                "code_version": _code_version(),
                "auto_eaf_path": str(eaf),
                "debug_json_path": str(eaf.with_name(eaf.stem.replace("_auto", "_debug") + ".json"))
                if eaf.with_name(eaf.stem.replace("_auto", "_debug") + ".json").exists()
                else "",
                "run_timestamp": _file_mtime_iso(eaf),
                "phase_unknown_count": meta.get("phase_unknown_count", meta.get("num_phase_unknowns", "")),
                "foot_unknown_count": meta.get("foot_unknown_count", meta.get("num_foot_unknowns", "")),
                "occlusion_count": meta.get("occlusion_count", ""),
                "validation_status": "unknown",
                "eaf_written": "TRUE",
                "exit_code": "",
                "error_message": "",
            }
    return rows


def _human_correction_rows(
    ground_truth_root: Path,
    video_rows: dict[str, dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    rows: dict[str, dict[str, Any]] = {}
    if not ground_truth_root.exists():
        return rows
    video_ids = set(video_rows)
    for eaf in sorted(ground_truth_root.rglob("*ground_truth*.eaf")):
        video_id = _match_video_id(eaf.stem, video_ids) or _safe_stem(eaf.stem.replace("_ground_truth", ""))
        correction_id = f"{video_id}_ground_truth"
        rows[correction_id] = {
            "correction_id": correction_id,
            "video_id": video_id,
            "auto_run_id": "",
            "reviewer": "",
            "auto_eaf_path": "",
            "ground_truth_eaf_path": str(eaf),
            "corrected_at": _file_mtime_iso(eaf),
            "review_time_minutes": "",
            "correction_status": "corrected",
            "notes": "",
        }
    return rows


def _pipeline_status_row(video_id: str, folder: Path, test_type: str) -> dict[str, Any]:
    biomarker = _first_existing(folder, ["biomarker.json", "**/biomarker.json"])
    sppb = _first_existing(folder, ["sppb_results.json", "**/sppb_results.json"])
    left = _first_boundary(folder, "left")
    right = _first_boundary(folder, "right")
    meta = _first_existing(folder, ["out2.csv", "out2.pkl", "meta*.csv", "meta*.pkl", "**/out2.csv", "**/out2.pkl"])
    tug_timing_found = bool(biomarker or sppb or meta)
    gait_events_found = bool(left or right or meta)
    known_artifacts = [item for item in (biomarker, sppb, left, right, meta) if item]
    if not known_artifacts:
        status = "missing"
    elif tug_timing_found and (_does_test_need_gait_events(test_type) is False or gait_events_found):
        status = "complete"
    else:
        status = "partial"
    return {
        "video_id": video_id,
        "pipeline_out_folder": str(folder),
        "biomarker_json": str(biomarker) if biomarker else "",
        "sppb_results_json": str(sppb) if sppb else "",
        "gait_boundaries_left": str(left) if left else "",
        "gait_boundaries_right": str(right) if right else "",
        "tug_timing_found": "TRUE" if tug_timing_found else "FALSE",
        "gait_events_found": "TRUE" if gait_events_found else "FALSE",
        "pipeline_status": status,
        "last_checked_at": _now_iso(),
        "error_message": "",
    }


def _is_visfrailty_v7_workbook(path: Path) -> bool:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return False
    try:
        return VISFRAILTY_V7_ANNOTATION_SHEET in wb.sheetnames
    finally:
        close = getattr(wb, "close", None)
        if close:
            close()


def _check_visfrailty_v7_outputs(
    excel_path: Path,
    visfrailty_root: str | Path | None,
) -> dict[str, int]:
    wb = load_workbook(excel_path)
    ws = wb[VISFRAILTY_V7_ANNOTATION_SHEET]
    headers = _v7_header_map(ws, VISFRAILTY_V7_HEADER_ROW)
    subjects = _v7_subject_rows(wb)
    root = Path(visfrailty_root) if visfrailty_root else None
    checked = 0
    complete = 0
    partial = 0
    missing = 0

    for row_idx in range(VISFRAILTY_V7_HEADER_ROW + 1, ws.max_row + 1):
        subject_id = str(_cell_value(ws, row_idx, headers, "subject_id") or "").upper()
        video_slot = str(_cell_value(ws, row_idx, headers, "video_slot") or "").upper()
        if not subject_id or not video_slot:
            continue
        subject = subjects.get(subject_id, {})
        output_run = Path(str(subject.get("output_run_path") or "")) if subject.get("output_run_path") else None
        candidate = _resolve_v7_output_folder(subject_id, video_slot, output_run, root)
        state = str(subject.get("processing_state") or "").strip().lower()
        has_output = bool(candidate and _has_pipeline_artifacts(candidate))
        if has_output or state in {"processed", "processed_no_json"}:
            status = "processed_no_json" if state == "processed_no_json" and not has_output else "processed"
            complete += 1
        elif candidate and candidate.exists():
            status = "partial"
            partial += 1
        else:
            status = "missing"
            missing += 1
        _set_cell(ws, row_idx, headers, "pipeline_state", status)
        _set_cell(ws, row_idx, headers, "biomarker_json", "yes" if has_output or status == "processed" else "")
        if candidate:
            _set_cell(ws, row_idx, headers, "pipeline_error", "" if status == "processed" else f"VisFrailTy output incomplete: {candidate}")
        _set_cell(ws, row_idx, headers, "flywheel_stage", _v7_stage_for_row(ws, row_idx, headers))
        checked += 1

    _refresh_visfrailty_v7_dashboard(wb)
    wb.save(excel_path)
    return {
        "videos_checked": checked,
        "pipeline_complete": complete,
        "pipeline_partial": partial,
        "pipeline_missing": missing,
    }


def _resolve_visfrailty_output_folder(
    input_folder: Path,
    video_id: str,
    test_type: str,
    existing_folder: Path | None,
    visfrailty_root: Path | None,
) -> Path:
    candidates: list[Path] = []
    if existing_folder and str(existing_folder) not in {"", "."}:
        candidates.append(existing_folder)
    candidates.extend(_visfrailty_output_dirs(input_folder, visfrailty_root))
    candidates.append(input_folder)

    seen: set[str] = set()
    unique = []
    for candidate in candidates:
        key = str(candidate)
        if key not in seen:
            unique.append(candidate)
            seen.add(key)
    for candidate in unique:
        if _has_pipeline_artifacts(candidate):
            return candidate
    return unique[0] if unique else input_folder


def _has_pipeline_artifacts(folder: Path) -> bool:
    return bool(
        _first_existing(
            folder,
            [
                "biomarker.json",
                "sppb_results.json",
                "out2.csv",
                "out2.pkl",
                "meta*.csv",
                "meta*.pkl",
                "**/biomarker.json",
                "**/sppb_results.json",
                "**/out2.csv",
                "**/out2.pkl",
            ],
        )
        or _first_boundary(folder, "left")
        or _first_boundary(folder, "right")
    )


def _visfrailty_output_dirs(folder: Path, explicit_root: Path | None = None) -> list[Path]:
    target_name = _normalized_name(folder.name)
    patient_id = _infer_patient_id(folder)
    if not target_name or not patient_id:
        return []

    candidates: list[Path] = []
    for output_root in _visfrailty_output_roots(folder, explicit_root):
        for patient_dir in _visfrailty_patient_output_dirs(output_root, patient_id):
            direct = patient_dir / folder.name
            if direct.is_dir():
                candidates.append(direct)
            try:
                children = list(patient_dir.iterdir())
            except OSError:
                continue
            for child in children:
                if child.is_dir() and _normalized_name(child.name) == target_name:
                    candidates.append(child)
    return sorted(set(candidates), key=lambda path: str(path).lower(), reverse=True)


def _visfrailty_output_roots(folder: Path, explicit_root: Path | None) -> list[Path]:
    roots: list[Path] = []
    if explicit_root:
        roots.extend(
            [
                explicit_root,
                explicit_root / "visfrailty_outputs",
                explicit_root / "output",
                explicit_root / "outputs",
            ]
        )
    inferred = _infer_volume_root(folder)
    if inferred is not None:
        roots.append(inferred / "visfrailty_outputs")
    roots.extend([folder.parent / "visfrailty_outputs", folder.parent.parent / "visfrailty_outputs"])
    return [root for root in _unique_paths(roots) if root.is_dir()]


def _visfrailty_patient_output_dirs(output_root: Path, patient_id: str) -> list[Path]:
    cache_key = (str(output_root), patient_id.lower())
    cached = _VISFRAILTY_PATIENT_OUTPUT_CACHE.get(cache_key)
    if cached is not None:
        return cached

    matches: list[Path] = []
    if patient_id.lower() in output_root.name.lower():
        matches.append(output_root)
    try:
        date_dirs = [path for path in output_root.iterdir() if path.is_dir()]
    except OSError:
        date_dirs = []
    for date_dir in date_dirs:
        if patient_id.lower() in date_dir.name.lower():
            matches.append(date_dir)
        try:
            children = [path for path in date_dir.iterdir() if path.is_dir()]
        except OSError:
            continue
        for child in children:
            if patient_id.lower() in child.name.lower():
                matches.append(child)

    matches = sorted(set(matches), key=lambda path: str(path).lower(), reverse=True)
    _VISFRAILTY_PATIENT_OUTPUT_CACHE[cache_key] = matches
    return matches


def _resolve_v7_output_folder(
    subject_id: str,
    video_slot: str,
    workbook_output_run: Path | None,
    explicit_root: Path | None,
) -> Path | None:
    roots: list[Path] = []
    if workbook_output_run:
        roots.append(workbook_output_run)
    if explicit_root:
        roots.extend([explicit_root, explicit_root / "visfrailty_outputs", explicit_root / "output", explicit_root / "outputs"])
    for root in _unique_paths(roots):
        if not root.exists():
            continue
        exact = root / f"{subject_id} {video_slot}"
        if exact.exists():
            return exact
        target = _normalized_name(f"{subject_id} {video_slot}")
        try:
            children = list(root.iterdir())
        except OSError:
            continue
        for child in children:
            if child.is_dir() and _normalized_name(child.name) == target:
                return child
    return workbook_output_run


def _visfrailty_v7_status(path: Path) -> WorkbookSummary:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[VISFRAILTY_V7_ANNOTATION_SHEET]
    headers = _v7_header_map(ws, VISFRAILTY_V7_HEADER_ROW)
    rows = list(_iter_v7_rows(ws, headers, VISFRAILTY_V7_HEADER_ROW))
    stages = [str(row.get("flywheel_stage") or _v7_stage_from_values(row)) for row in rows]
    auto = sum(1 for row in rows if _yes(row.get("auto_eaf_generated")) or row.get("auto_eaf_path"))
    corrected = sum(1 for row in rows if str(row.get("human_review_status") or "") not in {"", "not_started"} or row.get("human_eaf_path"))
    confirmed = sum(1 for row in rows if _yes(row.get("ground_truth_confirmed")))
    training = sum(1 for row in rows if _yes(row.get("usable_for_training")))
    raw = sum(1 for stage in stages if stage == "raw")
    pipeline_or_later = sum(1 for stage in stages if stage in VISFRAILTY_V7_STAGES[1:])
    auto_or_later = sum(1 for stage in stages if stage in VISFRAILTY_V7_STAGES[2:])
    return WorkbookSummary(
        videos=len(rows),
        pipeline_complete=pipeline_or_later,
        pipeline_partial=0,
        pipeline_missing=raw,
        auto_annotated=max(auto, auto_or_later),
        human_corrected=corrected,
        ground_truth_validated=confirmed,
        calibration_sample_videos=training,
        ready_for_auto_annotation=max(pipeline_or_later - auto, 0),
        ready_for_human_review=max(auto - corrected, 0),
        ready_for_calibration=max(confirmed - training, 0),
        review_queue=max(raw, 0) + max(auto - corrected, 0),
    )


def _v7_auto_annotation_map(
    output_root: Path,
    batch_output_root: Path,
) -> dict[tuple[str, str], dict[str, Any]]:
    rows = _auto_annotation_rows(output_root, batch_output_root)
    mapped: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows.values():
        path = Path(str(row.get("auto_eaf_path") or ""))
        key = _subject_slot_key_from_text(" ".join([path.name, str(path.parent), str(row.get("video_id") or "")]))
        if key:
            mapped[key] = row
    return mapped


def _v7_ground_truth_map(ground_truth_root: Path) -> dict[tuple[str, str], str]:
    mapped: dict[tuple[str, str], str] = {}
    if not ground_truth_root.exists():
        return mapped
    for eaf in sorted(ground_truth_root.rglob("*.eaf")):
        key = _subject_slot_key_from_text(" ".join([eaf.name, str(eaf.parent)]))
        if key:
            mapped[key] = str(eaf)
    return mapped


def _v7_unknown_count(auto_row: dict[str, Any]) -> int | str:
    phase = auto_row.get("phase_unknown_count")
    foot = auto_row.get("foot_unknown_count")
    if phase in {None, ""} and foot in {None, ""}:
        return ""
    return _int_or(phase, 0) + _int_or(foot, 0)


def _subject_slot_key_from_text(text: str) -> tuple[str, str] | None:
    upper = text.upper()
    subject = None
    subject_match = re.search(r"(^|[^A-Z0-9])((?:MC|NH|SG)\d{4}|NHGPAMK\d+|[A-Z]{2,}\d{4,})([^A-Z0-9]|$)", upper)
    if subject_match:
        subject = subject_match.group(2)
    slot = None
    for candidate in ("TUG1", "TUG2", "GS1", "GS2", "CRT", "SBS", "ST", "FT"):
        if re.search(rf"(^|[^A-Z0-9]){candidate}([^A-Z0-9]|$)", upper):
            slot = candidate
            break
    if subject and slot:
        return subject, slot
    return None


def _v7_subject_rows(wb: Any) -> dict[str, dict[str, Any]]:
    if VISFRAILTY_V7_SUBJECTS_SHEET not in wb.sheetnames:
        return {}
    ws = wb[VISFRAILTY_V7_SUBJECTS_SHEET]
    headers = {str(cell.value or ""): idx for idx, cell in enumerate(ws[1], start=1)}
    rows: dict[str, dict[str, Any]] = {}
    for row_idx in range(2, ws.max_row + 1):
        subject_id = ws.cell(row=row_idx, column=headers.get("subject_id", 0)).value if "subject_id" in headers else None
        if not subject_id:
            continue
        rows[str(subject_id).upper()] = {
            header: ws.cell(row=row_idx, column=col).value
            for header, col in headers.items()
            if header
        }
    return rows


def _refresh_visfrailty_v7_dashboard(wb: Any) -> None:
    if VISFRAILTY_V7_DASHBOARD_SHEET not in wb.sheetnames:
        return
    ws = wb[VISFRAILTY_V7_ANNOTATION_SHEET]
    headers = _v7_header_map(ws, VISFRAILTY_V7_HEADER_ROW)
    rows = list(_iter_v7_rows(ws, headers, VISFRAILTY_V7_HEADER_ROW))
    total = len(rows)
    counts = {stage: 0 for stage in VISFRAILTY_V7_STAGES}
    for row in rows:
        stage = str(row.get("flywheel_stage") or _v7_stage_from_values(row))
        if stage in counts:
            counts[stage] += 1
    dashboard = wb[VISFRAILTY_V7_DASHBOARD_SHEET]
    stage_rows = {
        "raw": 5,
        "pipeline_processed": 6,
        "auto_annotated": 7,
        "human_reviewed": 8,
        "ground_truth_confirmed": 9,
    }
    for stage, row_idx in stage_rows.items():
        count = counts.get(stage, 0)
        dashboard.cell(row=row_idx, column=2).value = count
        dashboard.cell(row=row_idx, column=3).value = f"{(count / total * 100):.1f}%" if total else "0.0%"
    dashboard.cell(row=10, column=2).value = total
    dashboard.cell(row=10, column=3).value = "100%" if total else "0%"


def _count_v7_annotation_rows(ws: Any) -> int:
    headers = _v7_header_map(ws, VISFRAILTY_V7_HEADER_ROW)
    return sum(1 for _ in _iter_v7_rows(ws, headers, VISFRAILTY_V7_HEADER_ROW))


def _iter_v7_rows(ws: Any, headers: dict[str, int], header_row: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    max_col = max(headers.values()) if headers else 0
    for values in ws.iter_rows(min_row=header_row + 1, max_col=max_col, values_only=True):
        row = {
            header: values[col - 1] if col <= len(values) else None
            for header, col in headers.items()
            if header
        }
        if row.get("subject_id") and row.get("video_slot"):
            rows.append(row)
    return rows


def _v7_header_map(ws: Any, header_row: int) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
    return {
        str(value or ""): idx
        for idx, value in enumerate(row, start=1)
        if value
    }


def _cell_value(ws: Any, row_idx: int, headers: dict[str, int], header: str) -> Any:
    col = headers.get(header)
    if col is None:
        return None
    return ws.cell(row=row_idx, column=col).value


def _set_cell(ws: Any, row_idx: int, headers: dict[str, int], header: str, value: Any) -> None:
    col = headers.get(header)
    if col is not None:
        ws.cell(row=row_idx, column=col).value = value


def _v7_stage_for_row(ws: Any, row_idx: int, headers: dict[str, int]) -> str:
    values = {header: _cell_value(ws, row_idx, headers, header) for header in headers}
    return _v7_stage_from_values(values)


def _v7_stage_from_values(values: dict[str, Any]) -> str:
    if _yes(values.get("usable_for_training")) or _yes(values.get("ground_truth_confirmed")):
        return "ground_truth_confirmed"
    review_status = str(values.get("human_review_status") or "").strip().lower()
    if review_status not in {"", "not_started"} or values.get("human_eaf_path"):
        return "human_reviewed"
    if _yes(values.get("auto_eaf_generated")) or values.get("auto_eaf_path"):
        return "auto_annotated"
    current_stage = str(values.get("flywheel_stage") or "").strip()
    if current_stage in VISFRAILTY_V7_STAGES:
        return current_stage
    pipeline_state = str(values.get("pipeline_state") or "").strip().lower()
    if pipeline_state in {"processed", "processed_no_json", "complete"} or _yes(values.get("biomarker_json")):
        return "pipeline_processed"
    return "raw"


def _yes(value: Any) -> bool:
    return str(value).strip().lower() in {"yes", "true", "1", "y"}


def _refresh_review_queue(wb: Any) -> None:
    videos = list(_iter_rows_as_dicts(wb["Videos"]))
    pipeline_by_video = {
        row.get("video_id"): row
        for row in _iter_rows_as_dicts(wb["PipelineOutputs"])
    }
    auto_by_video = {
        row.get("video_id"): row
        for row in _iter_rows_as_dicts(wb["AutoAnnotations"])
        if _truthy(row.get("eaf_written"))
    }
    corrected = {
        row.get("video_id")
        for row in _iter_rows_as_dicts(wb["HumanCorrections"])
        if str(row.get("correction_status") or "").lower() == "corrected"
    }
    rows: list[dict[str, Any]] = []
    for video in videos:
        video_id = video.get("video_id")
        if not video_id:
            continue
        pipeline = pipeline_by_video.get(video_id, {})
        auto = auto_by_video.get(video_id)
        pipeline_status = pipeline.get("pipeline_status") or "missing"
        if pipeline_status != "complete":
            rows.append(_review_row(10, video, auto, f"VisFrailTy {pipeline_status}", "run/check VisFrailTy output"))
        elif auto is None:
            rows.append(_review_row(20, video, auto, "missing auto annotation", "run auto annotation"))
        elif video_id not in corrected:
            rows.append(_review_row(30, video, auto, "missing human correction", "review in ELAN"))
        elif auto and str(auto.get("validation_status") or "") == "invalid":
            rows.append(_review_row(40, video, auto, "auto EAF invalid", "inspect debug output"))
    _write_plain_rows(wb["ReviewQueue"], rows)


def _review_row(priority: int, video: dict[str, Any], auto: dict[str, Any] | None, reason: str, action: str) -> dict[str, Any]:
    unknown = 0
    if auto:
        unknown = _int_or(auto.get("phase_unknown_count"), 0) + _int_or(auto.get("foot_unknown_count"), 0)
    return {
        "priority": priority,
        "video_id": video.get("video_id", ""),
        "test_type": video.get("test_type", ""),
        "auto_eaf_path": auto.get("auto_eaf_path", "") if auto else "",
        "reason": reason,
        "unknown_count": unknown,
        "validation_status": auto.get("validation_status", "") if auto else "",
        "suggested_action": action,
    }


def _update_video_statuses(wb: Any) -> None:
    videos = _sheet_rows_by_key(wb["Videos"], "video_id")
    pipeline = _sheet_rows_by_key(wb["PipelineOutputs"], "video_id")
    autos = _sheet_rows_by_key(wb["AutoAnnotations"], "video_id")
    corrections = _sheet_rows_by_key(wb["HumanCorrections"], "video_id")
    validations = _sheet_rows_by_key(wb["GroundTruthValidation"], "video_id")
    calibration_samples = {
        row.get("video_id")
        for row in _iter_rows_as_dicts(wb["CalibrationSamples"])
        if _truthy(row.get("included"))
    }
    for video_id, row in videos.items():
        if video_id in calibration_samples:
            row["status"] = "used_for_calibration"
        elif _truthy((validations.get(video_id) or {}).get("valid")):
            row["status"] = "validated"
        elif video_id in corrections:
            row["status"] = "human_corrected"
        elif video_id in autos:
            row["status"] = "auto_annotated"
        else:
            pipeline_status = (pipeline.get(video_id) or {}).get("pipeline_status")
            if pipeline_status == "complete":
                row["status"] = "visfrailty_complete"
            elif pipeline_status == "partial":
                row["status"] = "visfrailty_partial"
            elif pipeline_status == "missing":
                row["status"] = "visfrailty_missing"
            else:
                row.setdefault("status", "registered")
    _write_keyed_rows(wb["Videos"], "video_id", videos)


def _iter_rows_as_dicts(ws: Any) -> list[dict[str, Any]]:
    headers = _headers(ws)
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {headers[idx]: row[idx] for idx in range(min(len(headers), len(row))) if headers[idx]}
        if any(value not in {None, ""} for value in item.values()):
            rows.append(item)
    return rows


def _sheet_rows_by_key(ws: Any, key: str) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for row in _iter_rows_as_dicts(ws):
        value = row.get(key)
        if value not in {None, ""}:
            out[str(value)] = row
    return out


def _write_keyed_rows(ws: Any, key: str, rows: dict[str, dict[str, Any]]) -> None:
    headers = SHEET_HEADERS[ws.title]
    ws.delete_rows(2, max(ws.max_row - 1, 0))
    for row_idx, row in enumerate(sorted(rows.values(), key=lambda item: str(item.get(key) or "")), start=2):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx).value = row.get(header, "")


def _write_plain_rows(ws: Any, rows: list[dict[str, Any]]) -> None:
    headers = SHEET_HEADERS[ws.title]
    ws.delete_rows(2, max(ws.max_row - 1, 0))
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx).value = row.get(header, "")


def _headers(ws: Any) -> list[str]:
    return [str(cell.value or "") for cell in ws[1]]


def _column_for_header(ws: Any, header: str) -> str | None:
    for cell in ws[1]:
        if cell.value == header:
            return cell.column_letter
    return None


def _first_existing(folder: Path, patterns: list[str]) -> Path | None:
    if not folder.exists():
        return None
    for pattern in patterns:
        matches = sorted(folder.glob(pattern))
        if matches:
            return matches[0]
    return None


def _first_boundary(folder: Path, side: str) -> Path | None:
    if not folder.exists():
        return None
    side_re = re.compile(rf"(^|[^a-z]){side}([^a-z]|$)|(^|[^a-z]){side[0]}_?foot([^a-z]|$)", re.IGNORECASE)
    candidates = [
        path
        for path in folder.rglob("*.json")
        if "bound" in path.stem.lower()
        and side_re.search(path.stem)
        and "bounding_box" not in path.stem.lower()
    ]
    return sorted(candidates)[0] if candidates else None


def _does_test_need_gait_events(test_type: str) -> bool:
    return str(test_type or "").upper() in {"TUG", "GS1", "GS2", "GAIT_SPEED", "UNKNOWN"}


def _excel_test_type(value: Any) -> str:
    text = str(value or "").strip().lower()
    mapping = {
        "tug": "TUG",
        "crt": "CRT",
        "gait_speed": "GS1",
        "side_by_side": "SBS",
        "semi_tandem": "ST",
        "full_tandem": "FT",
        "balance": "unknown",
    }
    if text.upper() in TEST_TYPES:
        return text.upper()
    return mapping.get(text, "unknown")


def _video_id_from_auto_eaf(path: Path) -> str:
    stem = path.stem
    if stem.endswith("_auto"):
        stem = stem[:-5]
    return _safe_stem(stem)


def _match_video_id(stem: str, video_ids: set[str]) -> str | None:
    normalized = _safe_stem(stem.replace("_ground_truth", ""))
    if normalized in video_ids:
        return normalized
    for video_id in sorted(video_ids, key=len, reverse=True):
        if video_id and video_id in normalized:
            return video_id
    return None


def _infer_patient_id(folder: Path) -> str:
    for part in reversed(folder.parts):
        if re.search(r"[a-z]{2,}.*\d|\d{4,}", part.lower()):
            return part
    return folder.parent.name


def _infer_volume_root(folder: Path) -> Path | None:
    try:
        parts = folder.resolve().parts
    except OSError:
        parts = folder.parts
    if len(parts) >= 3 and parts[1] == "Volumes":
        return Path(*parts[:3])
    for parent in (folder, *folder.parents):
        if (parent / "visfrailty_outputs").is_dir():
            return parent
    return None


def _normalized_name(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name).lower())


def _unique_paths(paths: list[Path]) -> list[Path]:
    seen: set[str] = set()
    out: list[Path] = []
    for path in paths:
        key = str(path)
        if key not in seen:
            out.append(path)
            seen.add(key)
    return out


def _safe_relative(path: Path, root: Path) -> Path:
    try:
        return path.relative_to(root)
    except ValueError:
        return Path(path.name)


def _safe_stem(value: str) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    return text or "video"


def _unique_id(base: str, used: set[str]) -> str:
    candidate = base
    idx = 2
    while candidate in used:
        candidate = f"{base}_{idx}"
        idx += 1
    used.add(candidate)
    return candidate


def _read_json(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def _file_mtime_iso(path: Path) -> str:
    try:
        return datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()
    except OSError:
        return ""


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _code_version() -> str:
    head = Path(".git/HEAD")
    if not head.exists():
        return ""
    try:
        text = head.read_text(encoding="utf-8").strip()
        if text.startswith("ref: "):
            ref = Path(".git") / text[5:]
            return ref.read_text(encoding="utf-8").strip()[:12] if ref.exists() else ""
        return text[:12]
    except OSError:
        return ""


def _truthy(value: Any) -> bool:
    return str(value).strip().lower() in {"true", "1", "yes", "y"}


def _int_or(value: Any, default: int) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default
